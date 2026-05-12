"""Flask application for scraping YouTube video metadata and exporting to Excel."""

import io
import json
import os
import re
import sys
import urllib.request
from datetime import datetime

from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _resource_path(*parts: str) -> str:
    """Return the absolute path to a bundled resource (PyInstaller-aware)."""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, *parts)


app = Flask(
    __name__,
    template_folder=_resource_path("templates"),
    static_folder=_resource_path("static"),
)

_INNERTUBE_CONTEXT = {
    "client": {
        "clientName": "WEB",
        "clientVersion": "2.20231219.04.00",
        "hl": "en",
        "gl": "US",
    }
}

_HTTP_HEADERS = {
    "Content-Type": "application/json",
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/120.0.0.0 Safari/537.36"
    ),
}


def extract_video_id(url: str) -> str | None:
    """Return the YouTube video ID from a URL, or None if unparseable."""
    patterns = [
        r"(?:youtube\.com/watch\?v=|youtu\.be/|youtube\.com/embed/"
        r"|youtube\.com/v/|youtube\.com/shorts/)"
        r"([a-zA-Z0-9_-]{11})",
    ]
    for pattern in patterns:
        match = re.search(pattern, url.strip())
        if match:
            return match.group(1)
    stripped = url.strip()
    if re.fullmatch(r"[a-zA-Z0-9_-]{11}", stripped):
        return stripped
    return None


def _find_all(obj: dict | list, key: str) -> list:
    """Recursively find all values for *key* in a nested JSON structure."""
    results: list = []
    if isinstance(obj, dict):
        for k, v in obj.items():
            if k == key:
                results.append(v)
            results.extend(_find_all(v, key))
    elif isinstance(obj, list):
        for item in obj:
            results.extend(_find_all(item, key))
    return results


def _innertube_next(video_id: str) -> dict:
    """Call YouTube's innertube ``/next`` endpoint and return the JSON."""
    payload = json.dumps({"videoId": video_id, "context": _INNERTUBE_CONTEXT}).encode()
    req = urllib.request.Request(
        "https://www.youtube.com/youtubei/v1/next",
        data=payload,
        headers=_HTTP_HEADERS,
    )
    with urllib.request.urlopen(req, timeout=20) as resp:
        return json.loads(resp.read())


def _oembed(video_id: str) -> dict:
    """Call YouTube's oEmbed endpoint and return the JSON."""
    oembed_url = (
        f"https://www.youtube.com/oembed"
        f"?url=https://www.youtube.com/watch?v={video_id}&format=json"
    )
    req = urllib.request.Request(oembed_url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=10) as resp:
        return json.loads(resp.read())


def _parse_views(data: dict) -> int | None:
    """Extract the numeric view count from an innertube response."""
    for vcr in _find_all(data, "videoViewCountRenderer"):
        text = vcr.get("viewCount", {}).get("simpleText", "")
        digits = re.sub(r"[^0-9]", "", text)
        if digits:
            return int(digits)
    return None


def _parse_date(data: dict) -> str:
    """Extract and reformat the upload date as MM-DD-YYYY."""
    for dt in _find_all(data, "dateText"):
        text = dt.get("simpleText", "")
        if text:
            for fmt in ("%b %d, %Y", "%B %d, %Y"):
                try:
                    return datetime.strptime(text, fmt).strftime("%m-%d-%Y")
                except ValueError:
                    continue
            return text
    return "N/A"


def _parse_likes(data: dict) -> int | None:
    """Extract the like count from accessibility text."""
    for text in _find_all(data, "accessibilityText"):
        if isinstance(text, str) and "like this video along with" in text:
            match = re.search(r"along with ([\d,]+)", text)
            if match:
                return int(match.group(1).replace(",", ""))
    return None


def _parse_comments(data: dict) -> str | None:
    """Extract the comment count string from the engagement panel header."""
    for panel in _find_all(data, "engagementPanelTitleHeaderRenderer"):
        title_runs = panel.get("title", {}).get("runs", [])
        if any("Comment" in r.get("text", "") for r in title_runs):
            ctx_runs = panel.get("contextualInfo", {}).get("runs", [])
            if ctx_runs:
                return ctx_runs[0].get("text")
    return None


def fetch_video_metadata(url: str) -> dict:
    """Fetch metadata for a single YouTube video."""
    video_id = extract_video_id(url)
    if not video_id:
        return {"error": f"Invalid YouTube URL: {url}", "url": url}

    canonical_url = f"https://www.youtube.com/watch?v={video_id}"

    try:
        oembed_data = _oembed(video_id)
        title = oembed_data.get("title", "N/A")
        channel = oembed_data.get("author_name", "N/A")
    except Exception:
        title = "N/A"
        channel = "N/A"

    try:
        next_data = _innertube_next(video_id)
        views = _parse_views(next_data)
        formatted_date = _parse_date(next_data)
        likes = _parse_likes(next_data)
        comments = _parse_comments(next_data)
    except Exception as exc:
        return {
            "error": str(exc),
            "url": canonical_url,
        }

    return {
        "date": formatted_date,
        "channel": channel,
        "url": canonical_url,
        "title": title,
        "views": views,
        "likes": likes,
        "comments": comments,
    }


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/scrape", methods=["POST"])
def scrape():
    """Accept a list of YouTube URLs and return metadata for each."""
    data = request.get_json(silent=True) or {}
    urls = data.get("urls", [])

    if not urls:
        return jsonify({"error": "No URLs provided"}), 400

    results = []
    for url in urls:
        result = fetch_video_metadata(url)
        results.append(result)

    return jsonify({"results": results})


@app.route("/api/export", methods=["POST"])
def export():
    """Accept scraped results and return an .xlsx file."""
    data = request.get_json(silent=True) or {}
    results = data.get("results", [])

    if not results:
        return jsonify({"error": "No results to export"}), 400

    wb = Workbook()
    ws = wb.active
    ws.title = "YouTube Metadata"

    headers = [
        "Date (M-D-Y)",
        "Channel Name",
        "YouTube Link",
        "Video Title",
        "View Count",
        "Like Count",
        "Comment Count",
    ]

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row_idx, item in enumerate(results, 2):
        if "error" in item:
            ws.cell(row=row_idx, column=1, value="Error")
            ws.cell(row=row_idx, column=3, value=item.get("url", ""))
            ws.cell(row=row_idx, column=4, value=item.get("error", ""))
            continue

        ws.cell(row=row_idx, column=1, value=item.get("date", "N/A"))
        ws.cell(row=row_idx, column=2, value=item.get("channel", "N/A"))

        link_cell = ws.cell(row=row_idx, column=3, value=item.get("url", ""))
        link_cell.font = Font(color="0563C1", underline="single")

        ws.cell(row=row_idx, column=4, value=item.get("title", "N/A"))

        views = item.get("views")
        ws.cell(row=row_idx, column=5, value=views if views is not None else "N/A")

        likes = item.get("likes")
        ws.cell(row=row_idx, column=6, value=likes if likes is not None else "N/A")

        comments = item.get("comments")
        ws.cell(row=row_idx, column=7, value=comments if comments is not None else "N/A")

    column_widths = [14, 22, 45, 50, 14, 14, 14]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="youtube_metadata.xlsx",
    )


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
